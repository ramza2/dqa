#!/usr/bin/env python3
"""M2-2E synthetic-only DOCX/ZIP layout QA fixture generator.

Runs from backend/ with the development virtualenv. Never queries Oracle,
PostgreSQL, DEMIS or patient data. Output is ignored local QA artifacts.
The structural assertions are NOT a substitute for rendering every page.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
from dataclasses import replace
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile
import sys

# Support both `python -m scripts.m2_2e_docx_layout_qa` and direct script launch.
# Direct launch normally places only backend/scripts on sys.path, excluding tests/.
if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm

from app.services.schema_snapshot_report import build_snapshot_table_docx
from app.services.schema_snapshot_schema_bundle import build_snapshot_schema_zip
from tests.test_m2_2c_snapshot_report import _stored as table_fixture
from tests.test_m2_2d_schema_bundle import _stored as schema_fixture
from tests.test_schema_snapshot_history import FakeSession, _base, _capture


def _stress_row():
    snap = deepcopy(_base())
    snap.tables[0] = replace(
        snap.tables[0],
        table_comment="합성 자료입니다. " + "긴 테이블 설명과 한글 줄바꿈 확인. " * 10,
    )
    base_column = snap.columns[1]
    snap.columns = [
        replace(snap.columns[0], column_comment="기본키 설명 · 한글 가나다 abc 123"),
        replace(base_column, column_comment="한글·영문·숫자가 섞인 컬럼"),
    ]
    for i in range(3, 31):
        snap.columns.append(replace(
            base_column,
            ordinal_position=i,
            column_name="COL_" + str(i).zfill(2) + "_" + "LONG_IDENTIFIER_" * 5,
            column_comment=("한글 보고서 품질 확인: 열 너비·단어 줄바꿈·머리글 반복. " * 6),
        ))
    return _capture(FakeSession(), snapshot=snap).revision


def _doc_text(doc: Document) -> str:
    return "\n".join(
        [p.text for p in doc.paragraphs]
        + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    )


def _inspect_docx(payload: bytes, *, expected: tuple[str, ...]) -> dict:
    assert payload[:2] == b"PK", "not a valid OOXML ZIP header"
    doc = Document(BytesIO(payload))
    assert len(doc.sections) == 1
    sect = doc.sections[0]
    assert abs(sect.page_width - Cm(21)) <= Cm(.01), "not A4 width"
    assert abs(sect.page_height - Cm(29.7)) <= Cm(.01), "not A4 height"
    assert doc.tables, "no tables in analysis report"
    for table in doc.tables:
        assert table.rows, "an empty table"
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        assert tr_pr.find(qn("w:tblHeader")) is not None, "non-repeating table heading"
    content = _doc_text(doc)
    assert "NOT_VERIFIED" in content
    for token in expected:
        assert token in content, f"expected term omitted: {token}"
    for forbidden in ("SECRET_SQL_DEFAULT_NEVER_EXPORT",
                      "DANGEROUS_SQL_DEFAULT_SHOULD_NEVER_EXPORT",
                      "DANGEROUS_INDEX_SQL_SHOULD_NEVER_EXPORT"):
        assert forbidden not in content, "sensitive SQL entered the report"
    return {
        "size": len(payload),
        "sha256": sha256(payload).hexdigest(),
        "tables": len(doc.tables),
        "cells": sum(len(row.cells) for table in doc.tables for row in table.rows),
        "max_cell_characters": max(
            len(cell.text) for table in doc.tables for row in table.rows for cell in row.cells
        ),
        "headings": [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")],
        "a4_portrait": True,
        "all_table_headers_repeat": True,
        "rendered_pages": "NOT_VERIFIED_REQUIRES_RENDER",
    }


def generate(directory: Path) -> dict:
    directory.mkdir(parents=True, exist_ok=True)
    one = table_fixture()
    stress = _stress_row()
    one_file = directory / "single-table-standard.docx"
    stress_file = directory / "single-table-korean-stress.docx"
    one_file.write_bytes(build_snapshot_table_docx(
        one, schema_name="DEMIS_OWNER", table_name="TB_MOCK"
    ))
    stress_file.write_bytes(build_snapshot_table_docx(
        stress, schema_name="DEMIS_OWNER", table_name="TB_MOCK"
    ))
    documents = {
        one_file.name: _inspect_docx(one_file.read_bytes(), expected=(
            "PK_MOCK", "UK_MOCK", "IX_MOCK", "FK_MOCK_PARENT", "DEMIS_OWNER",
        )),
        stress_file.name: _inspect_docx(stress_file.read_bytes(), expected=(
            "한글", "LONG_IDENTIFIER", "DEMIS_OWNER",
        )),
    }

    batch = schema_fixture()
    zip_entries = []
    offset = 0
    while True:
        archive_bytes, total, next_offset = build_snapshot_schema_zip(
            batch, schema_name="DEMIS_OWNER", offset=offset, limit=2,
        )
        archive_file = directory / f"schema-batch-{offset:03d}.zip"
        archive_file.write_bytes(archive_bytes)
        with ZipFile(BytesIO(archive_bytes)) as archive:
            manifest = json.loads(archive.read("manifest.json"))
            assert manifest["source_name"] == batch.source_name
            assert manifest["snapshot_revision_id"] == batch.id
            assert manifest["schema_fingerprint"] == batch.schema_fingerprint
            assert manifest["total_tables_in_schema"] == total
            assert manifest["next_offset"] == next_offset
            files = manifest["files"]
            assert len(files) == manifest["exported_table_count"]
            for entry in files:
                part = archive.read(entry["file"])
                assert sha256(part).hexdigest() == entry["sha256"]
                exported_file = directory / f"batch-{entry['ordinal_in_schema']:03d}.docx"
                exported_file.write_bytes(part)
                documents[exported_file.name] = _inspect_docx(part, expected=(
                    "DEMIS_OWNER", entry["table_name"],
                ))
        zip_entries.append({
            "file": archive_file.name,
            "offset": offset,
            "next_offset": next_offset,
            "table_names": [x["table_name"] for x in files],
            "sha256": sha256(archive_bytes).hexdigest(),
        })
        if next_offset is None:
            break
        assert next_offset > offset
        offset = next_offset

    assert sum(len(x["table_names"]) for x in zip_entries) == total
    result = {
        "scope": "SYNTHETIC_METADATA_ONLY",
        "source": "M2-2C/M2-2D test fixtures",
        "structural_checks": "PASS",
        "visual_check": "PENDING_RENDER_AND_HUMAN_INSPECTION",
        "documents": documents,
        "zip_parts": zip_entries,
    }
    (directory / "qa-manifest.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8",
    )
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--out-dir", type=Path, default=Path("artifacts/m2-2e-docx-qa"),
        help="local-only generated files (no PII or DEMIS data)",
    )
    args = parser.parse_args()
    result = generate(args.out_dir)
    print("M2_2E_STRUCTURAL_QA_PASS")
    print("Output:", args.out_dir.resolve())
    print("DOCX:", ", ".join(result["documents"]))
    print("ZIP:", ", ".join(x["file"] for x in result["zip_parts"]))
    print("VISUAL_QA_NOT_YET_COMPLETE: render and inspect all DOCX pages")


if __name__ == "__main__":
    main()
