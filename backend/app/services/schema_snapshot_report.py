"""Bounded DOCX table analysis report from verified, stored Schema Snapshot.

This report is an *unapproved metadata preview*, not a DEMIS database read or
Catalog READY artifact. It does not export Oracle default SQL, index
definitions, target credentials, or any clinical data rows.
"""
from __future__ import annotations

from collections import defaultdict
from io import BytesIO
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.models.schema_snapshot_revision import SchemaSnapshotRevision
from app.services.schema_snapshot_erd import SnapshotErdError, build_snapshot_erd
from app.services.schema_snapshot_history import _snapshot_from_payload

_MAX_COLUMNS = 400
_MAX_FOREIGN_KEYS = 200
_MAX_INDEXES = 100
_MAX_DOCX_BYTES = 3_000_000
_INVALID_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F\uFFFE\uFFFF]")


class SnapshotReportError(Exception):
    def __init__(self, code: str):
        self.code = code
        super().__init__(code)


def _clean(value: object, limit: int = 600) -> str:
    if value is None:
        return "—"
    cleaned = _INVALID_XML.sub("", str(value))
    return cleaned[:limit] if cleaned else "—"


# A4 portrait inner width: 21.0 - 1.6 - 1.6 = 17.8 cm.
# Keep names readable, but allocate the most space to actual DB COMMENTS.
_COLUMN_DETAIL_WIDTHS_CM = (0.65, 4.05, 2.45, 1.75, 1.15, 7.75)


def _grid(
    document: Document,
    headers: list[str],
    rows: list[list[object]],
    *,
    column_widths_cm: tuple[float, ...] | None = None,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, item in enumerate(headers):
        table.rows[0].cells[i].text = item
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = _clean(value)
    if column_widths_cm is not None:
        # The default equal-width, auto-fit layout makes comments a tiny column
        # when long identifiers force Word to expand the Column cell.
        if len(column_widths_cm) != len(headers):
            raise ValueError("report column widths must match headers")
        table.autofit = False
        for column, width in zip(table.columns, column_widths_cm):
            column.width = Cm(width)
        for row in table.rows:
            for cell, width in zip(row.cells, column_widths_cm):
                cell.width = Cm(width)
        # Compact margins and header lettering avoid wrapping Nullable / Keys.
        cell_margins = table._tbl.tblPr.find(qn("w:tblCellMar"))
        if cell_margins is None:
            cell_margins = OxmlElement("w:tblCellMar")
            table._tbl.tblPr.append(cell_margins)
        for side in ("start", "end"):
            margin = cell_margins.find(qn("w:" + side))
            if margin is None:
                margin = OxmlElement("w:" + side)
                cell_margins.append(margin)
            margin.set(qn("w:w"), "57")  # approx. 0.1 cm in twips
            margin.set(qn("w:type"), "dxa")
        for cell in table.rows[0].cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    # Keep the column names visible when the table spans multiple pages.
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    tr_pr.append(repeat_header)


def _configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.6)
    style = document.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style.font.size = Pt(9)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    document.core_properties.author = "DQA"
    document.core_properties.last_modified_by = "DQA"
    document.core_properties.title = "DQA DB 물리구조 분석서 (미승인 Snapshot)"
    document.core_properties.subject = "Snapshot metadata only"


def build_snapshot_table_docx(
    row: SchemaSnapshotRevision, *, schema_name: str, table_name: str,
) -> bytes:
    """Produce one focused table's physical structure; never full raw Snapshot."""
    try:
        # M2-2B validates stored v1 structure, target, fingerprint and counts.
        build_snapshot_erd(row, focus_schema=schema_name, focus_table=table_name)
        snapshot = _snapshot_from_payload(row.snapshot_json)
        selected = next(
            t for t in snapshot.tables
            if t.schema_name == schema_name and t.table_name == table_name
        )
    except SnapshotErdError as exc:
        code = ("SNAPSHOT_REPORT_TABLE_NOT_FOUND"
                if exc.code == "SNAPSHOT_ERD_TABLE_NOT_FOUND"
                else "SNAPSHOT_REPORT_UNVERIFIABLE")
        raise SnapshotReportError(code) from None
    except (StopIteration, ValueError, KeyError, TypeError, AttributeError):
        raise SnapshotReportError("SNAPSHOT_REPORT_UNVERIFIABLE") from None

    columns = sorted(
        (c for c in snapshot.columns
         if c.schema_name == schema_name and c.table_name == table_name),
        key=lambda c: (c.ordinal_position, c.column_name),
    )
    pk = sorted(
        (c for c in snapshot.primary_keys
         if c.schema_name == schema_name and c.table_name == table_name),
        key=lambda c: (c.constraint_name, c.ordinal_position),
    )
    unique = sorted(
        (c for c in snapshot.unique_constraints
         if c.schema_name == schema_name and c.table_name == table_name),
        key=lambda c: (c.constraint_name, c.ordinal_position),
    )
    outbound = [
        fk for fk in snapshot.foreign_keys
        if fk.schema_name == schema_name and fk.source_table == table_name
    ]
    inbound = [
        fk for fk in snapshot.foreign_keys
        if fk.target_schema == schema_name and fk.target_table == table_name
    ]
    indexes = sorted(
        (ix for ix in snapshot.indexes
         if ix.schema_name == schema_name and ix.table_name == table_name),
        key=lambda ix: ix.index_name,
    )
    if (len(columns) > _MAX_COLUMNS or len(outbound) + len(inbound) > _MAX_FOREIGN_KEYS
            or len(indexes) > _MAX_INDEXES):
        raise SnapshotReportError("SNAPSHOT_REPORT_SCOPE_TOO_LARGE")

    doc = Document()
    _configure(doc)
    title = doc.add_heading("DEMIS DB 물리구조 분석서", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    disclaimer = doc.add_paragraph(
        "검토용 · 미승인 Snapshot 메타데이터 · 독립 DBA Preflight NOT_VERIFIED · "
        "Catalog READY/활성화 및 DEMIS 데이터 조회 권한을 의미하지 않음"
    )
    disclaimer.style = doc.styles["Subtitle"]
    doc.add_paragraph("분석 범위: 선택 테이블 1개 (전체 DB 분석서가 아님)")
    doc.add_heading("1. 분석 대상 및 출처", 1)
    _grid(doc, ["항목", "값"], [
        ["Source", row.source_name],
        ["DBMS", row.db_type],
        ["Database", row.database_name],
        ["Schema", schema_name],
        ["Table", table_name],
        ["Snapshot Revision", row.id],
        ["Schema Fingerprint", row.schema_fingerprint],
        ["Snapshot Capture", row.captured_at.isoformat() if row.captured_at else "—"],
        ["검증 상태", "NOT_VERIFIED (독립 DBA 승인 아님)"],
    ])
    doc.add_paragraph(
        "물리 구조는 저장된 Snapshot에서 확인한 항목만 사용합니다. "
        "DB COMMENT는 원문 메타데이터로 표시하며 추정 업무 의미를 생성하지 않습니다. "
        "계정·접속정보·컬럼 기본값(SQL)·인덱스 정의 SQL·의료데이터 원문은 제외합니다."
    )
    doc.add_heading("2. 테이블 개요", 1)
    _grid(doc, ["항목", "값"], [
        ["Schema.Table", schema_name + "." + table_name],
        ["Table Type", selected.table_type],
        ["Table DB COMMENT", selected.table_comment],
        ["Columns", len(columns)],
        ["Outbound / Inbound FK", str(len(outbound)) + " / " + str(len(inbound))],
        ["Indexes", len(indexes)],
    ])

    doc.add_heading("3. 컬럼 상세", 1)
    keys = {(x.column_name) for x in pk}
    uniques = {(x.column_name) for x in unique}
    _grid(doc, ["#", "Column", "Type", "Nullable", "Keys", "DB COMMENT"], [
        [
            c.ordinal_position, c.column_name,
            (c.data_type + ("(" + str(c.character_maximum_length) + ")"
                            if c.character_maximum_length is not None else
                            ("(" + str(c.numeric_precision) +
                             ("," + str(c.numeric_scale) if c.numeric_scale is not None else "") + ")"
                             if c.numeric_precision is not None else ""))),
            "Y" if c.is_nullable else "N",
            ", ".join(v for v in ("PK" if c.column_name in keys else "",
                                  "UK 구성원" if c.column_name in uniques else "") if v) or "—",
            c.column_comment,
        ]
        for c in columns
    ], column_widths_cm=_COLUMN_DETAIL_WIDTHS_CM)
    doc.add_heading("4. PK 및 UK 구성", 1)
    grouped: dict[tuple[str, str], list[object]] = defaultdict(list)
    for c in pk:
        grouped[("PK", c.constraint_name)].append(c.column_name)
    for c in unique:
        grouped[("UK", c.constraint_name)].append(c.column_name)
    if grouped:
        _grid(doc, ["구분", "Constraint", "구성 컬럼 (순서 유지)"], [
            [kind, constraint, ", ".join(names)]
            for (kind, constraint), names in sorted(grouped.items())
        ])
    else:
        doc.add_paragraph("Snapshot에 PK/UK 메타데이터가 없습니다.")

    doc.add_heading("5. FK 관계", 1)
    relations = []
    for direction, group in (("OUT", outbound), ("IN", inbound)):
        for fk in group:
            mapping = ", ".join(
                c.source_column + " → " + c.target_column
                for c in sorted(fk.columns, key=lambda c: c.ordinal_position)
            )
            relations.append([
                direction, fk.constraint_name,
                fk.schema_name + "." + fk.source_table,
                fk.target_schema + "." + fk.target_table,
                mapping,
            ])
    if relations:
        _grid(doc, ["방향", "Constraint", "Source", "Target", "컬럼 매핑"], relations)
        doc.add_paragraph("외부 Schema/Table 참조는 존재 및 조회 권한이 확인되지 않은 참조 이름입니다.")
    else:
        doc.add_paragraph("Snapshot에 해당 테이블의 FK 메타데이터가 없습니다.")

    doc.add_heading("6. 인덱스", 1)
    if indexes:
        _grid(doc, ["Index", "Unique", "Method", "Columns"], [
            [ix.index_name, "Y" if ix.is_unique else "N", ix.index_method,
             ", ".join(c.column_name for c in sorted(ix.columns, key=lambda c: c.ordinal_position))]
            for ix in indexes
        ])
    else:
        doc.add_paragraph("Snapshot에 인덱스 메타데이터가 없습니다.")

    doc.add_paragraph(
        "이 문서는 저장된 물리 메타데이터의 조회 결과이며, 실제 의료정보 매핑·"
        "DBA 권한 승인·FHIR 데이터 항목 계약을 확정하지 않습니다."
    )
    output = BytesIO()
    try:
        doc.save(output)
    except Exception:
        raise SnapshotReportError("SNAPSHOT_REPORT_RENDER_FAILED") from None
    content = output.getvalue()
    if len(content) > _MAX_DOCX_BYTES:
        raise SnapshotReportError("SNAPSHOT_REPORT_SCOPE_TOO_LARGE")
    return content
