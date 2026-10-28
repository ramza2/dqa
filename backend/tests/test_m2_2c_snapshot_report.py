"""M2-2C: synthetic Snapshot DOCX, zero external DB / Oracle use."""
from copy import deepcopy
from dataclasses import replace
from datetime import datetime, timezone
from io import BytesIO

from docx import Document
import pytest

from app.adapters.schema_inspector.base import (
    InspectedForeignKey, InspectedForeignKeyColumn,
    InspectedIndex, InspectedIndexColumn, InspectedPrimaryKey,
    InspectedUniqueConstraint,
)
from app.services.schema_snapshot_report import (
    SnapshotReportError, build_snapshot_table_docx,
)
from tests.test_schema_snapshot_history import FakeSession, _base, _capture

URL = "/api/v1/schema-inspector/snapshots/1/report.docx"


def _stored():
    snapshot = deepcopy(_base())
    snapshot.columns[0] = replace(
        snapshot.columns[0], default_value="DANGEROUS_SQL_DEFAULT_SHOULD_NEVER_EXPORT",
    )
    snapshot.primary_keys = [
        InspectedPrimaryKey("DEMIS_OWNER", "TB_MOCK", "PK_MOCK", "ID", 1)
    ]
    snapshot.unique_constraints = [
        InspectedUniqueConstraint("DEMIS_OWNER", "TB_MOCK", "UK_MOCK", "CODE", 1)
    ]
    snapshot.foreign_keys = [
        InspectedForeignKey(
            schema_name="DEMIS_OWNER", constraint_name="FK_MOCK_PARENT",
            source_table="TB_MOCK", target_schema="REF_OWNER",
            target_table="TB_PARENT",
            columns=(
                InspectedForeignKeyColumn(2, "CODE", "SITE"),
                InspectedForeignKeyColumn(1, "ID", "PARENT_ID"),
            ),
        )
    ]
    snapshot.indexes = [
        InspectedIndex(
            schema_name="DEMIS_OWNER", table_name="TB_MOCK",
            index_name="IX_MOCK", is_unique=False, index_method="NORMAL",
            index_definition="DANGEROUS_INDEX_SQL_SHOULD_NEVER_EXPORT",
            columns=(InspectedIndexColumn(1, "CODE"),),
        )
    ]
    db = FakeSession()
    row = _capture(db, snapshot=snapshot).revision
    row.captured_at = datetime(2026, 10, 28, tzinfo=timezone.utc)
    return row


def _all_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    return "\n".join(
        [p.text for p in document.paragraphs]
        + [cell.text for table in document.tables
           for row in table.rows for cell in row.cells]
    )


def test_docx_has_identity_columns_keys_foreign_key_and_index_without_sql():
    row = _stored()
    content = build_snapshot_table_docx(
        row, schema_name="DEMIS_OWNER", table_name="TB_MOCK",
    )
    assert content.startswith(b"PK")
    text = _all_text(content)
    for item in [
        "DEMIS_OWNER", "TB_MOCK", "PK_MOCK", "UK_MOCK", "IX_MOCK",
        "FK_MOCK_PARENT", "PARENT_ID", "CODE → SITE", "NOT_VERIFIED",
        row.schema_fingerprint, "선택 테이블 1개",
    ]:
        assert item in text
    for secret in [
        "DANGEROUS_SQL_DEFAULT_SHOULD_NEVER_EXPORT",
        "DANGEROUS_INDEX_SQL_SHOULD_NEVER_EXPORT",
        "capture_request_id", "password", "index_definition",
    ]:
        assert secret not in text
        assert secret.encode() not in content
    assert len(content) <= 3_000_000


@pytest.mark.parametrize("damage", ["fingerprint", "counts", "version", "payload", "identity"])
def test_corruption_rejected_before_document_export(damage):
    row = _stored()
    if damage == "fingerprint":
        row.schema_fingerprint = "a" * 64
    elif damage == "counts":
        row.counts_json["columns"] = 999
    elif damage == "version":
        row.snapshot_version = 100
    elif damage == "payload":
        row.snapshot_json["tables"][0]["table_comment"] = "PRIVATE_MUTATION"
    elif damage == "identity":
        row.schema_name = "WRONG_SCHEMA"
    with pytest.raises(SnapshotReportError, match="SNAPSHOT_REPORT_UNVERIFIABLE"):
        build_snapshot_table_docx(row, schema_name="DEMIS_OWNER", table_name="TB_MOCK")


def test_unknown_table_fails_without_document():
    row = _stored()
    with pytest.raises(SnapshotReportError, match="SNAPSHOT_REPORT_TABLE_NOT_FOUND"):
        build_snapshot_table_docx(row, schema_name="OTHER_SCHEMA", table_name="TB_MOCK")


def test_report_api_requires_manage_permission_and_is_source_scoped(
    unauth_client, client, monkeypatch,
):
    from app.api.routes import schema_snapshot_history as api
    row = _stored()
    monkeypatch.setattr(api, "get_snapshot_revision",
                        lambda _session, revision_id, *, source_name:
                        row if revision_id == 1 and source_name == row.source_name else None)
    params = {"source_name": row.source_name, "schema_name": "DEMIS_OWNER", "table_name": "TB_MOCK"}
    assert unauth_client.get(URL, params=params).status_code == 401
    viewer = unauth_client.get(
        URL, params=params,
        headers={"X-DQA-Dev-Actor": "reader", "X-DQA-Dev-Roles": "viewer"},
    )
    assert viewer.status_code == 403

    response = client.get(URL, params=params)
    assert response.status_code == 200, response.text[:200]
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment;" in response.headers["content-disposition"]
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert "FK_MOCK_PARENT" in _all_text(response.content)
    assert client.get(URL, params={**params, "source_name": "other_source"}).status_code == 404
    assert client.get(URL, params={**params, "table_name": "UNKNOWN"}).status_code == 404


def test_corrupt_report_endpoint_returns_sanitized_409(client, monkeypatch):
    from app.api.routes import schema_snapshot_history as api
    row = _stored()
    row.snapshot_json["tables"][0]["table_comment"] = "HIGHLY_PRIVATE_VALUE"
    monkeypatch.setattr(api, "get_snapshot_revision",
                        lambda _session, revision_id, *, source_name: row)
    res = client.get(URL, params={
        "source_name": row.source_name, "schema_name": "DEMIS_OWNER", "table_name": "TB_MOCK",
    })
    assert res.status_code == 409
    assert res.json()["detail"]["code"] == "SNAPSHOT_REPORT_UNVERIFIABLE"
    assert "HIGHLY_PRIVATE_VALUE" not in res.text
