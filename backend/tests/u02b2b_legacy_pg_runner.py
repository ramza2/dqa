"""Original Legacy finalizer with real isolated PostgreSQL Catalog persistence.

Runs inside U-02-B-2-B Compose with pinned original Legacy checkout mounted
read-only. Reuses the original integration test's _create_fixture/_cleanup
which write 2 analysis runs, 2 physical snapshots, real Catalog/Preflight
rows into a disposable database. No Oracle or medical DB connection exists.
"""
from __future__ import annotations

import io
import json
import sys
import zipfile
from pathlib import Path

from docx import Document
from sqlalchemy import select, func

from app.db.catalog_bootstrap import ensure_catalog_schema
from app.db.session import get_catalog_session_factory
from app.models.catalog import CatalogAnalysisRun, CatalogSource, CatalogTable
from app.models.catalog_category import CatalogCategory, CatalogTableCategory
from app.models.catalog_history import CatalogAnalysisSnapshot
from app.models.catalog_validation import CatalogPreflightResult
from app.services.catalog_package_finalizer import build_final_catalog_package
from tests.test_catalog_package_finalization import _create_fixture, _cleanup


OUTPUT = Path("/output")
EXPECTED_DB = "u02b2b_catalog"
EXPECTED_HOST = "catalog-db"
STATUS_VALUES = ("READY", "WARNING", "BLOCKED")


def assert_isolation() -> None:
    import os

    assert os.environ["CATALOG_DB_HOST"] == EXPECTED_HOST
    assert os.environ["CATALOG_DB_NAME"] == EXPECTED_DB
    assert os.environ["CATALOG_DB_USER"] == "u02b2b_user"
    assert os.environ["MEDICAL_DB_HOST"] == "non-existent-test-medical-host"
    assert Path("/legacy/backend/tests/test_catalog_package_finalization.py").is_file()
    assert OUTPUT.is_dir()


def persist_category(source_id: int) -> None:
    with get_catalog_session_factory()() as db:
        table = db.scalar(
            select(CatalogTable).where(
                CatalogTable.source_id == source_id,
                CatalogTable.table_name == "TB_PT_MST",
            )
        )
        assert table is not None
        category = CatalogCategory(
            source_id=source_id, category_key="U02B2B_TEST",
            category_name="Synthetic-only category", description="Integration fixture",
            sort_order=1, active=True,
        )
        db.add(category)
        db.flush()
        db.add(CatalogTableCategory(
            table_id=table.id, category_id=category.id,
            is_primary=True, assignment_source="MANUAL",
            confidence=1.0, note="fixture only",
        ))
        db.commit()


def persisted_evidence(source_id: int) -> None:
    with get_catalog_session_factory()() as db:
        assert db.scalar(select(func.count()).select_from(CatalogAnalysisRun).where(
            CatalogAnalysisRun.source_id == source_id)) == 2
        assert db.scalar(select(func.count()).select_from(CatalogAnalysisSnapshot).where(
            CatalogAnalysisSnapshot.source_id == source_id)) == 2
        assert db.scalar(select(func.count()).select_from(CatalogPreflightResult).where(
            CatalogPreflightResult.source_id == source_id)) == 1
        assert db.scalar(select(func.count()).select_from(CatalogCategory).where(
            CatalogCategory.source_id == source_id)) == 1


def set_preflight_status(source_id: int, status: str) -> None:
    assert status in STATUS_VALUES
    with get_catalog_session_factory()() as db:
        row = db.scalar(select(CatalogPreflightResult).where(
            CatalogPreflightResult.source_id == source_id))
        assert row is not None
        # PostgreSQL JSONB must receive a new dict for ORM dirty tracking.
        row.result_json = {**row.result_json, "status": status}
        row.status = status
        db.commit()


def emit_legacy_zip(source_id: int, status: str) -> dict:
    with get_catalog_session_factory()() as db:
        result = build_final_catalog_package(db, source_id)
    assert result.manifest["package_readiness"] == status
    assert result.manifest["package_version"] == "2.0"
    artifact = OUTPUT / f"final_{status}.zip"
    artifact.write_bytes(result.content)

    with zipfile.ZipFile(io.BytesIO(result.content)) as outer:
        prefix = "demis_catalog_package/"
        manifest = json.loads(outer.read(prefix + "manifest.json"))
        names = set(outer.namelist())
        reports = [x for x in names if x.startswith(prefix + "reports/") and x.endswith(".docx")]
        assert len(reports) == 1
        report_bytes = outer.read(reports[0])
        # Actual python-docx renderer; never a fake .docx placeholder.
        report = Document(io.BytesIO(report_bytes))
        report_text = "\n".join(
            [p.text for p in report.paragraphs]
            + [cell.text for tbl in report.tables for row in tbl.rows for cell in row.cells]
        )
        assert "TB_PT_MST" in report_text or "PT_ID" in report_text
        assert manifest["artifacts"]["db_analysis_report"]["available"] is True
        assert manifest["artifacts"]["latest_diff"]["available"] is True
        assert manifest["artifacts"]["latest_diff"]["identical"] is True
        assert manifest["artifacts"]["preflight"]["status"] == status
        diff = json.loads(outer.read(prefix + "diff/latest.json"))
        assert diff["available"] is True
        assert diff["identical"] is True
        assert diff["base_run_id"] != diff["target_run_id"]
    return {
        "status": status,
        "filename": artifact.name,
        "file_count": len(manifest["files"]),
        "report_verified_as_docx": True,
        "diff_verified": True,
    }


def main() -> None:
    assert_isolation()
    ensure_catalog_schema()
    source_id = None
    try:
        source_id = _create_fixture()
        persist_category(source_id)
        persisted_evidence(source_id)
        output = []
        for status in STATUS_VALUES:
            set_preflight_status(source_id, status)
            output.append(emit_legacy_zip(source_id, status))
        (OUTPUT / "generation_summary.json").write_text(
            json.dumps({"fixture": "u02b2b-disposable-legacy-pg", "results": output}, indent=2),
            encoding="utf-8",
        )
        print(json.dumps({"status": "PASS", "generated": [x["filename"] for x in output]}))
    finally:
        if source_id is not None:
            _cleanup(source_id)


if __name__ == "__main__":
    main()
