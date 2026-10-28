"""M2-2E synthetic report structure QA without DB or LibreOffice."""
import json
from io import BytesIO
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

import pytest

from docx import Document

from scripts.m2_2e_docx_layout_qa import generate


def test_qa_generator_produces_standard_stress_and_complete_batch(tmp_path):
    result = generate(tmp_path)
    assert result["scope"] == "SYNTHETIC_METADATA_ONLY"
    assert result["structural_checks"] == "PASS"
    assert result["visual_check"] == "PENDING_RENDER_AND_HUMAN_INSPECTION"
    assert "single-table-standard.docx" in result["documents"]
    assert "single-table-korean-stress.docx" in result["documents"]
    assert len(result["zip_parts"]) == 2
    assert [part["offset"] for part in result["zip_parts"]] == [0, 2]
    assert len(result["documents"]) == 5  # two singles + three ZIP exports

    manifest = json.loads((tmp_path / "qa-manifest.json").read_text(encoding="utf-8"))
    assert manifest["documents"] == result["documents"]
    assert manifest["zip_parts"] == result["zip_parts"]

    doc = Document(tmp_path / "single-table-korean-stress.docx")
    assert len(doc.tables) >= 2
    assert "NOT_VERIFIED" in "\n".join(p.text for p in doc.paragraphs)
    cells = [c.text for table in doc.tables for row in table.rows for c in row.cells]
    assert any("LONG_IDENTIFIER" in text for text in cells)
    assert any("한글" in text for text in cells)
    widths = [round(column.width.cm, 2) for column in doc.tables[2].columns]
    assert widths == [0.65, 4.05, 2.45, 1.75, 1.15, 7.75]
    assert doc.tables[2].autofit is False
    assert result["documents"]["single-table-korean-stress.docx"]["column_detail_widths_cm"] == widths
    assert widths[5] > widths[1]
    for batch in result["zip_parts"]:
        with ZipFile(tmp_path / batch["file"]) as archive:
            assert "manifest.json" in archive.namelist()
            assert "README.txt" in archive.namelist()
            for entry in json.loads(archive.read("manifest.json"))["files"]:
                content = archive.read(entry["file"])
                assert content.startswith(b"PK")
                assert Document(BytesIO(content)).tables


@pytest.mark.parametrize("entrypoint", ["module", "file"])
def test_qa_cli_entrypoints_find_backend_tests_package(tmp_path, entrypoint):
    """A subprocess catches the sys.path bug that regular pytest masked."""
    backend_root = Path(__file__).resolve().parents[1]
    target = tmp_path / entrypoint
    command = (
        [sys.executable, "-m", "scripts.m2_2e_docx_layout_qa"]
        if entrypoint == "module"
        else [sys.executable, "scripts/m2_2e_docx_layout_qa.py"]
    )
    completed = subprocess.run(
        command + ["--out-dir", str(target)],
        cwd=backend_root,
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert completed.returncode == 0, completed.stdout + "\n" + completed.stderr
    assert "M2_2E_STRUCTURAL_QA_PASS" in completed.stdout
    manifest = json.loads((target / "qa-manifest.json").read_text(encoding="utf-8"))
    assert manifest["structural_checks"] == "PASS"
    assert manifest["visual_check"] == "PENDING_RENDER_AND_HUMAN_INSPECTION"
    assert (target / "single-table-korean-stress.docx").is_file()
